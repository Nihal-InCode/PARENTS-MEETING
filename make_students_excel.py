import sqlite3
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB = "web_attendance (11).db"
OUT = "students_by_class.xlsx"

conn = sqlite3.connect(DB)
rows = conn.execute(
    "SELECT id, roll_no, name, class FROM students ORDER BY class, CAST(roll_no AS INTEGER)"
).fetchall()
conn.close()

classes = defaultdict(list)
for sid, roll, name, cls in rows:
    classes[cls or "UNCLASSIFIED"].append((sid, roll, name))

thin = Side(style="thin", color="999999")
box_border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(bold=True, color="FFFFFF")
title_font = Font(bold=True, size=14)
fills = ["FFF2CC", "E2EFDA", "FCE4D6", "DDEBF7", "E4DFEC", "F2F2F2", "D9E1F2", "EDEDED", "FDE9D9", "DFEEDF"]


def style_header(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center")
    c.border = box_border
    return c


wb = Workbook()
ws = wb.active
ws.title = "All Classes"

ws.cell(row=1, column=1, value="Students - Name & ID by Class").font = title_font

col = 1
for i, cls in enumerate(sorted(classes)):
    fill = PatternFill("solid", fgColor=fills[i % len(fills)])
    start_col = col

    t = ws.cell(row=3, column=col, value=cls)
    t.font = Font(bold=True, size=12)
    t.fill = fill
    t.alignment = Alignment(horizontal="center")
    t.border = box_border
    ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=start_col + 2)
    for cc in range(start_col, start_col + 3):
        ws.cell(row=3, column=cc).border = box_border
        ws.cell(row=3, column=cc).fill = fill

    style_header(ws, 4, col, "Student ID")
    style_header(ws, 4, col + 1, "Roll No")
    style_header(ws, 4, col + 2, "Name")

    r = 5
    for sid, roll, name in classes[cls]:
        vals = (sid, roll, name)
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=col + j, value=v)
            c.border = box_border
            c.fill = fill
            if j == 0:
                c.alignment = Alignment(horizontal="center")
        r += 1

    ws.column_dimensions[get_column_letter(start_col)].width = 12
    ws.column_dimensions[get_column_letter(start_col + 1)].width = 10
    ws.column_dimensions[get_column_letter(start_col + 2)].width = 28
    col += 4

for cls in sorted(classes):
    s = wb.create_sheet(cls)
    style_header(s, 1, 1, "Student ID")
    style_header(s, 1, 2, "Roll No")
    style_header(s, 1, 3, "Name")
    for i, (sid, roll, name) in enumerate(classes[cls], start=2):
        s.cell(row=i, column=1, value=sid).alignment = Alignment(horizontal="center")
        s.cell(row=i, column=2, value=roll).alignment = Alignment(horizontal="center")
        s.cell(row=i, column=3, value=name)
        for c in range(1, 4):
            s.cell(row=i, column=c).border = box_border
    s.column_dimensions["A"].width = 12
    s.column_dimensions["B"].width = 10
    s.column_dimensions["C"].width = 30
    s.freeze_panes = "A2"

wb.save(OUT)
print("Saved", OUT, "with", sum(len(v) for v in classes.values()), "students in", len(classes), "classes")
